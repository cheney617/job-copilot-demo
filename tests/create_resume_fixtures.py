from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "tmp" / "pdfs"
TEXT = [
    "Synthetic Resume - Senior Product Manager",
    "Led consumer growth, subscriptions, retention, paid conversion and revenue initiatives.",
    "Owned a 0-to-1 product launch from user problem discovery through MVP and iteration.",
    "Built AI Agent workflows using LLM capabilities and intent recognition.",
    "Designed A/B testing and analytics metrics to validate outcomes.",
    "Coordinated cross-functional global teams across six functions.",
]


def create_pdf() -> Path:
    path = OUT / "job-copilot-test-resume.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    text = pdf.beginText(54, 790)
    text.setFont("Helvetica", 12)
    for line in TEXT:
        text.textLine(line)
    pdf.drawText(text)
    pdf.save()
    return path


def create_docx() -> Path:
    path = OUT / "job-copilot-test-resume.docx"
    document = Document()
    document.add_heading(TEXT[0], level=1)
    for line in TEXT[1:]:
        document.add_paragraph(line)
    document.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(create_pdf())
    print(create_docx())
